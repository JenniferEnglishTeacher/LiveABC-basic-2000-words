import json, re, shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

root = Path(__file__).resolve().parents[2]
unit_dir = root / "unit-26"
data = json.loads((unit_dir / "unit-26-content.json").read_text(encoding="utf-8"))
source_root = root.parents[1]
shutil.copy2(source_root / "Track 51.mp3", unit_dir / "Track 51.mp3")
shutil.copy2(source_root / "Track 52.mp3", unit_dir / "Track 52.mp3")

template = (root / "unit-04/unit-04-quiz.html").read_text(encoding="utf-8")
html = re.sub(r"<title>[\s\S]*?</title>", "<title>Protect Yourself from Blue Light Vocabulary &amp; Reading Quiz</title>", template, count=1)
html = re.sub(r"<h1>[\s\S]*?</h1>", "<h1>Protect Yourself from Blue Light</h1>", html, count=1)
html = html.replace("Reading — Track 07", "Reading — Track 51").replace('src="Track 07.mp3"', 'src="Track 51.mp3"')
html = html.replace("Vocabulary — Track 08", "Vocabulary — Track 52").replace('src="Track 08.mp3"', 'src="Track 52.mp3"')
embedded = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
html = re.sub(r'<script id="unit-data" type="application/json">[\s\S]*?</script>', lambda _: f'<script id="unit-data" type="application/json">{embedded}</script>', html, count=1)
(unit_dir / "unit-26-quiz.html").write_text(html, encoding="utf-8")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55)
sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
title = doc.add_heading("LiveABC Basic 2000 Words — Unit 26", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(data["title"]).bold = True
doc.add_paragraph("Name: ____________________    Class: __________    Date: __________")
doc.add_heading("Part 1 — Listening Cloze", level=1)
doc.add_paragraph("Listen to Track 51 and fill in the 24 missing words.")
cloze = data["cloze"]["text"]
for a in data["cloze"]["answers"]:
    cloze = cloze.replace("{{" + str(a["blank"]) + "}}", f'({a["blank"]}) ' + "_" * 14)
doc.add_paragraph(cloze)
doc.add_heading("Part 2 — Reading Comprehension", level=1)
letters = "ABCD"
for i, q in enumerate(data["reading_questions"], 1):
    doc.add_paragraph(f"{i}. {q['question']}")
    doc.add_paragraph("    " + "    ".join(f"({letters[j]}) {choice}" for j, choice in enumerate(q["choices"])))
doc.add_heading("Part 3 — Vocabulary in Context", level=1)
doc.add_paragraph("Word bank: " + " • ".join(item["word"] for item in data["vocabulary"]))
for i, item in enumerate(data["vocabulary"], 1): doc.add_paragraph(f"{i}. {item['quiz_sentence']}")
doc.add_page_break()
doc.add_heading("Teacher Answer Key — Unit 26", 0)
doc.add_heading("Part 1", level=1); doc.add_paragraph("; ".join(f'{a["blank"]}. {a["answer"]}' for a in data["cloze"]["answers"]))
doc.add_heading("Part 2", level=1); doc.add_paragraph("; ".join(f'{i}. {letters[q["answer"]]}' for i, q in enumerate(data["reading_questions"], 1)))
doc.add_heading("Part 3", level=1); doc.add_paragraph("; ".join(f'{i}. {item["word"]}' for i, item in enumerate(data["vocabulary"], 1)))
doc.add_heading("Part 4 — Spell and Read Aloud", level=1)
doc.add_paragraph("Students spell the target form used in each sample sentence. The browser pronunciation check is optional.")
table = doc.add_table(rows=12, cols=2)
for i, item in enumerate(data["vocabulary"], 1):
    cell = table.cell((i - 1) % 12, (i - 1) // 12)
    run = cell.paragraphs[0].add_run(f'{i}. {item["word"]}: {item["sample_sentence"]}')
    run.font.size = Pt(8)
doc.save(unit_dir / "unit-26-test.docx")

sent_path = root / "scripts/speech-sentences.json"
sent = json.loads(sent_path.read_text(encoding="utf-8"))
sent["26"] = [item["sample_sentence"] for item in data["vocabulary"]]
sent_path.write_text(json.dumps(sent, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

index_path = root / "index.html"
index = index_path.read_text(encoding="utf-8")
if "UNIT 26" not in index:
    card = '    <section class="card"><div class="tag">UNIT 26</div><h2>Protect Yourself from Blue Light</h2><p>Listening, reading, vocabulary, and target-word spelling and optional read-aloud practice.</p><a class="button" href="unit-26/unit-26-quiz.html">Start Unit 26</a><a class="download" href="unit-26/unit-26-quiz.html#speech">Spell & read aloud · 拼字與朗讀</a><a class="download" href="unit-26/unit-26-test.docx">Download Word test</a></section>\n'
    index = index.replace('  </div>\n  <p class="note">', card + '  </div>\n  <p class="note">')
    index_path.write_text(index, encoding="utf-8")

readme_path = root / "README.md"
readme = readme_path.read_text(encoding="utf-8")
if "Unit 26 —" not in readme:
    readme = readme.replace("- [Unit 25 — Vampires and Dracula](unit-25/unit-25-quiz.html)", "- [Unit 25 — Vampires and Dracula](unit-25/unit-25-quiz.html)\n- [Unit 26 — Protect Yourself from Blue Light](unit-26/unit-26-quiz.html)")
readme = readme.replace("Units 1–25", "Units 1–26")
readme_path.write_text(readme, encoding="utf-8")

for rel in ("scripts/validate-speaking.cjs", "scripts/browser-test.cjs"):
    path = root / rel
    text = path.read_text(encoding="utf-8").replace("n<=25", "n<=26")
    if rel.endswith("browser-test.cjs"):
        text = text.replace("count(),25", "count(),26")
    if rel.endswith("validate-speaking.cjs"):
        text = text.replace("Validated 25 units", "Validated 26 units")
    path.write_text(text, encoding="utf-8")
