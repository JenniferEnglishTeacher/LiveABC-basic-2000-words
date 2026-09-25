import json, re, shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

root = Path(__file__).resolve().parents[2]
unit_dir = root / "unit-21"
data = json.loads((unit_dir / "unit-21-content.json").read_text(encoding="utf-8"))
source_root = root.parents[1]
shutil.copy2(source_root / "Track 41.mp3", unit_dir / "Track 41.mp3")
shutil.copy2(source_root / "Track 42.mp3", unit_dir / "Track 42.mp3")

template = (root / "unit-04/unit-04-quiz.html").read_text(encoding="utf-8")
html = re.sub(r"<title>[\s\S]*?</title>", "<title>Mr. Johnson&#39;s Greetings from Spain Vocabulary &amp; Reading Quiz</title>", template, count=1)
html = re.sub(r"<h1>[\s\S]*?</h1>", "<h1>Mr. Johnson&#39;s Greetings from Spain</h1>", html, count=1)
html = html.replace("Reading — Track 07", "Reading — Track 41").replace('src="Track 07.mp3"', 'src="Track 41.mp3"')
html = html.replace("Vocabulary — Track 08", "Vocabulary — Track 42").replace('src="Track 08.mp3"', 'src="Track 42.mp3"')
embedded = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
html = re.sub(r'<script id="unit-data" type="application/json">[\s\S]*?</script>', lambda _: f'<script id="unit-data" type="application/json">{embedded}</script>', html, count=1)
(unit_dir / "unit-21-quiz.html").write_text(html, encoding="utf-8")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55)
sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
title = doc.add_heading("LiveABC Basic 2000 Words — Unit 21", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(data["title"]).bold = True
doc.add_paragraph("Name: ____________________    Class: __________    Date: __________")
doc.add_heading("Part 1 — Listening Cloze", level=1)
doc.add_paragraph("Listen to Track 41 and fill in the 21 missing words.")
cloze = data["cloze"]["text"]
for a in data["cloze"]["answers"]:
    cloze = cloze.replace("{{" + str(a["blank"]) + "}}", f'({a["blank"]}) ' + "_" * 14)
doc.add_paragraph(cloze)
doc.add_heading("Part 2 — Reading Comprehension", level=1)
letters = "ABCD"
for i, q in enumerate(data["reading_questions"], 1):
    doc.add_paragraph(f"{i}. {q['question']}")
    doc.add_paragraph("    " + "    ".join(f"({letters[j]}) {choice}" for j, choice in enumerate(q["choices"])))
doc.add_heading("Part 3 — Vocabulary in Context", level=1)
doc.add_paragraph("Word bank: " + " • ".join(item["word"] for item in data["vocabulary"]))
for i, item in enumerate(data["vocabulary"], 1): doc.add_paragraph(f"{i}. {item['quiz_sentence']}")
doc.add_page_break()
doc.add_heading("Teacher Answer Key — Unit 21", 0)
doc.add_heading("Part 1", level=1); doc.add_paragraph("; ".join(f'{a["blank"]}. {a["answer"]}' for a in data["cloze"]["answers"]))
doc.add_heading("Part 2", level=1); doc.add_paragraph("; ".join(f'{i}. {letters[q["answer"]]}' for i, q in enumerate(data["reading_questions"], 1)))
doc.add_heading("Part 3", level=1); doc.add_paragraph("; ".join(f'{i}. {item["word"]}' for i, item in enumerate(data["vocabulary"], 1)))
doc.add_heading("Part 4 — Spell and Read Aloud", level=1)
doc.add_paragraph("Students spell the target form used in each sample sentence. The browser pronunciation check is optional.")
table = doc.add_table(rows=11, cols=2)
for i, item in enumerate(data["vocabulary"], 1):
    cell = table.cell((i - 1) % 11, (i - 1) // 11)
    run = cell.paragraphs[0].add_run(f'{i}. {item["word"]}: {item["sample_sentence"]}')
    run.font.size = Pt(8)
doc.save(unit_dir / "unit-21-test.docx")

sent_path = root / "scripts/speech-sentences.json"
sent = json.loads(sent_path.read_text(encoding="utf-8"))
sent["21"] = [item["sample_sentence"] for item in data["vocabulary"]]
sent_path.write_text(json.dumps(sent, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

index_path = root / "index.html"
index = index_path.read_text(encoding="utf-8")
if "UNIT 21" not in index:
    card = '    <section class="card"><div class="tag">UNIT 21</div><h2>Mr. Johnson&#39;s Greetings from Spain</h2><p>Listening, reading, vocabulary, and target-word spelling and optional read-aloud practice.</p><a class="button" href="unit-21/unit-21-quiz.html">Start Unit 21</a><a class="download" href="unit-21/unit-21-quiz.html#speech">Spell & read aloud · 拼字與朗讀</a><a class="download" href="unit-21/unit-21-test.docx">Download Word test</a></section>\n'
    index = index.replace('  </div>\n  <p class="note">', card + '  </div>\n  <p class="note">')
    index_path.write_text(index, encoding="utf-8")

readme_path = root / "README.md"
readme = readme_path.read_text(encoding="utf-8")
if "Unit 21 —" not in readme:
    readme = readme.replace("- [Unit 20 — Rosefield Tour Bus](unit-20/unit-20-quiz.html)", "- [Unit 20 — Rosefield Tour Bus](unit-20/unit-20-quiz.html)\n- [Unit 21 — Mr. Johnson's Greetings from Spain](unit-21/unit-21-quiz.html)")
readme = readme.replace("Units 1–20", "Units 1–21")
readme_path.write_text(readme, encoding="utf-8")

for rel in ("scripts/validate-speaking.cjs", "scripts/browser-test.cjs"):
    path = root / rel
    text = path.read_text(encoding="utf-8").replace("n<=20", "n<=21")
    if rel.endswith("browser-test.cjs"):
        text = text.replace("count(),20", "count(),21")
    if rel.endswith("validate-speaking.cjs"):
        text = text.replace("Validated 20 units", "Validated 21 units")
    path.write_text(text, encoding="utf-8")
