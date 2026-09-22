import json
from pathlib import Path
from gtts import gTTS
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

root = Path(__file__).resolve().parents[2]
unit_dir = root / "unit-04"
data = json.loads((unit_dir / "unit-04-content.json").read_text(encoding="utf-8"))

gTTS(data["passage"], lang="en").save(str(unit_dir / "Track 07.mp3"))
vocab_audio = "Unit 4 vocabulary. " + " ".join(
    f'{i}. {item["word"]}. {item["sample_sentence"]}' for i, item in enumerate(data["vocabulary"], 1)
)
gTTS(vocab_audio, lang="en").save(str(unit_dir / "Track 08.mp3"))

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55)
sec.left_margin = Inches(.65); sec.right_margin = Inches(.65)
styles = doc.styles
styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(10)
title = doc.add_heading("LiveABC Basic 2000 Words — Unit 4", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(data["title"]).bold = True
doc.add_paragraph("Name: ____________________    Class: __________    Date: __________")

doc.add_heading("Part 1 — Listening Cloze", level=1)
doc.add_paragraph("Listen to Track 07 and fill in the 17 missing words or phrases.")
cloze = data["cloze"]["text"]
for answer in data["cloze"]["answers"]:
    cloze = cloze.replace("{{" + str(answer["blank"]) + "}}", f'({answer["blank"]}) ' + "_" * 14)
doc.add_paragraph(cloze)

doc.add_heading("Part 2 — Reading Comprehension", level=1)
letters = "ABCD"
for i, q in enumerate(data["reading_questions"], 1):
    doc.add_paragraph(f"{i}. {q['question']}")
    doc.add_paragraph("    " + "    ".join(f"({letters[j]}) {choice}" for j, choice in enumerate(q["choices"])))

doc.add_heading("Part 3 — Vocabulary in Context", level=1)
doc.add_paragraph("Word bank: " + " • ".join(item["word"] for item in data["vocabulary"]))
for i, item in enumerate(data["vocabulary"], 1):
    doc.add_paragraph(f"{i}. {item['quiz_sentence']}")

doc.add_page_break()
doc.add_heading("Teacher Answer Key — Unit 4", 0)
doc.add_heading("Part 1", level=1)
doc.add_paragraph("; ".join(f'{a["blank"]}. {a["answer"]}' for a in data["cloze"]["answers"]))
doc.add_heading("Part 2", level=1)
doc.add_paragraph("; ".join(f'{i}. {letters[q["answer"]]}' for i, q in enumerate(data["reading_questions"], 1)))
doc.add_heading("Part 3", level=1)
doc.add_paragraph("; ".join(f'{i}. {item["word"]}' for i, item in enumerate(data["vocabulary"], 1)))
doc.add_heading("Part 4 — Spell and Read Aloud", level=1)
doc.add_paragraph("Students spell the target form used in each sample sentence. The browser pronunciation check is optional.")
for i, item in enumerate(data["vocabulary"], 1):
    doc.add_paragraph(f'{i}. {item["word"]}: {item["sample_sentence"]}')

doc.save(unit_dir / "unit-04-test.docx")
